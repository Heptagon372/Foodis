"""④ 한식진흥원 「한식메뉴 외국어표기 길라잡이 800선」 반영 (한국 음식 표기 표준화).

준비: https://www.data.go.kr/data/15129784/fileData.do 에서 XLSX 다운로드 → data/raw/hansik800.xlsx
라이선스: 공공데이터, 이용허락범위 제한 없음 (출처 표기 권장)
출력: data/raw/hansik800.json {slug: {name_ko, name_en_std, desc_ko, desc_en}}
열 이름이 배포 회차마다 다를 수 있어 헤더를 키워드로 찾는다.
"""
from __future__ import annotations

import re

from common import RAW, SEED, read_csv, write_json

HEADER_HINTS = {
    "name_ko": ["한글", "메뉴명", "국문"],
    "name_en": ["영문", "영어"],
    "desc_ko": ["설명", "국문설명", "한글설명"],
    "desc_en": ["영문설명", "영어설명", "description"],
}


def find_cols(header: list[str]) -> dict[str, int]:
    cols: dict[str, int] = {}
    norm = [re.sub(r"\s", "", str(h or "")).lower() for h in header]
    for key, hints in HEADER_HINTS.items():
        for i, h in enumerate(norm):
            if any(x.lower() in h for x in hints) and i not in cols.values():
                if key == "name_en" and "설명" in h:
                    continue
                cols[key] = i
                break
    return cols


def norm_ko(s: str) -> str:
    return re.sub(r"[\s\-·()]", "", s or "")


def main() -> None:
    from openpyxl import load_workbook
    path = RAW / "hansik800.xlsx"
    if not path.exists():
        raise SystemExit("data/raw/hansik800.xlsx 가 없습니다. 공공데이터포털에서 내려받아 넣어 주세요.")
    ws = load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows[:10]) if sum(1 for c in r if c) >= 3)
    cols = find_cols(list(rows[hi]))
    print("감지된 열:", {k: rows[hi][v] for k, v in cols.items()})
    table = {norm_ko(str(r[cols["name_ko"]])): r for r in rows[hi + 1:] if r and r[cols["name_ko"]]}
    out = {}
    for t in read_csv(SEED / "dish_targets.csv"):
        if t["country_code"] != "KR":
            continue
        r = table.get(norm_ko(t["name_ko"]))
        if r:
            g = lambda k: (str(r[cols[k]]).strip() if k in cols and r[cols[k]] else None)  # noqa: E731
            out[t["slug"]] = {"name_ko": g("name_ko"), "name_en_std": g("name_en"),
                              "desc_ko": g("desc_ko"), "desc_en": g("desc_en")}
    write_json(RAW / "hansik800.json", out)
    print(f"✔ 한국 음식 {len(out)}건 표기 표준 매칭 → data/raw/hansik800.json")


if __name__ == "__main__":
    main()
